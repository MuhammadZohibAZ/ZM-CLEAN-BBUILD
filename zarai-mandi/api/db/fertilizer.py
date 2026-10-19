"""Fertilizer prices (data/Fast uni.xlsx), loaded by etl_sqlite.py.

The fertilizer export is separate from the commodity export:
  - sheet 1 ("fertilizer cleaned formate"): Date, Province, Station,
    Input_Company, Input_Product_En, Arrivalsunit, Minimum, Maximum
  - sheet 2 ("fertlizer company, brand, prod"): Company, Brand, Products

How it is loaded (agreed with the product owner):
  - one by-product per brand product, named "<Brand> <Product>" (the data
    already says "Sona Urea", "Engro DAP"; a name without its brand, such as
    FMC's "Ammonium Sulphate", becomes "FMC Ammonium Sulphate");
  - the rows have no location, so they are a national price: province
    "Pakistan", district "Pakistan", station "All Pakistan";
  - the rows have no rate type: "Dealer Rate";
  - prices are per 50 kg bag (Arrivalsunit = 50; the app labels the unit);
  - only days inside the commodity export's date range are kept, so the app's
    "latest day" (the newest date in the database) stays the commodity one;
  - the generic fertilizer catalog entries (UREA, DAP, MOP, ...) are replaced
    by the brand products; pesticide entries in the Fertilizer division stay.

It also writes app/src/customer-face/shared/data/fertilizerBrands.json
(by-product -> company, brand), which the by-product cards show as "Company".
"""
import json
import os
import re

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.join(BASE_DIR, "..", "..", "data", "Fast uni.xlsx")
APP_BRANDS_JSON = os.path.join(
    BASE_DIR, "..", "..", "app", "src", "customer-face", "shared", "data", "fertilizerBrands.json"
)

DIVISION = "Fertilizer"
RATE_TYPE = "Dealer Rate"
PROVINCE, DISTRICT, STATION = "Pakistan", "Pakistan", "All Pakistan"
ROW_NUM_BASE = 10_000_000  # row_num is unique; commodity rows start at 1

# Generic catalog entries the brand products replace (pesticides are kept).
GENERIC_FERTILIZERS = {
    "Ammonium Sulphate", "Amonium Nitrate", "CAN", "DAP", "Enrich", "Green Phosphate",
    "MOP", "NP", "NPK", "Pak Arab CAN", "Pak Arab Guara", "SOP-G", "SSP", "TSP", "UREA",
    "Zabardast Urea", "Zarkhez", "Zinc", "Zinc Sulphate",
}


def _s(v):
    return str(v).strip() if v not in (None, "") else ""


def _key(s):
    return re.sub(r"[^a-z0-9]", "", s.lower())


def read_brands(wb):
    """[(company, brand, [products])] from the company / brand sheet."""
    ws = wb.worksheets[1]
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        company, brand, products = (_s(row[0]), _s(row[1]), _s(row[2])) if row else ("", "", "")
        if company and brand:
            out.append((company, brand, [p.strip() for p in products.split(",") if p.strip()]))
    return out


def brand_for(company, product, brands):
    """The brand of one product: the company's brand whose name starts the
    product, else the one listing it, else the company's only brand."""
    own = [b for b in brands if b[0] == company]
    for _, brand, _ in own:
        if _key(product).startswith(_key(brand)):
            return brand
    for _, brand, products in own:
        if any(_key(p) == _key(product) or _key(product).endswith(_key(p)) for p in products):
            return brand
    return own[0][1] if len(own) == 1 else None


def display_name(product, brand):
    if brand and not _key(product).startswith(_key(brand)):
        return f"{brand} {product}"
    return product


def read(xlsx_path=DEFAULT_XLSX):
    """Rows ready for price_records + {by_product: {company, brand}}."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    brands = read_brands(wb)
    ws = wb.worksheets[0]
    rows, info, skipped = [], {}, 0
    head = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if head is None:
            head = {h: k for k, h in enumerate(row) if h}
            continue
        date_v = row[head["Date"]]
        company = _s(row[head["Input_Company"]])
        product = _s(row[head["Input_Product_En"]])
        lo, hi = row[head["Minimum"]], row[head["Maximum"]]
        if not date_v or not company or not product or lo in (None, "") or hi in (None, ""):
            skipped += 1
            continue
        brand = brand_for(company, product, brands)
        name = display_name(product, brand)
        info.setdefault(name, {"company": company, "brand": brand or ""})
        record_date = date_v.strftime("%Y-%m-%d") if hasattr(date_v, "strftime") else str(date_v).split(" ")[0]
        unit = row[head["Arrivalsunit"]] if "Arrivalsunit" in head else None
        rows.append((record_date, name, float(lo), float(hi), float(unit) if unit not in (None, "") else None))
    return rows, info, skipped


def load(cur, xlsx_path=DEFAULT_XLSX):
    """Replace the generic fertilizer catalog entries with the brand products and
    insert their prices. Call after the catalog CSV and the commodity rows are in."""
    rows, info, skipped = read(xlsx_path)
    cur.execute("SELECT min(record_date), max(record_date) FROM price_records")
    first, last = cur.fetchone()
    trimmed = 0
    if first and last:
        kept = [r for r in rows if first <= r[0] <= last]
        trimmed = len(rows) - len(kept)
        rows = kept
    names = sorted({r[1] for r in rows})
    info = {n: info[n] for n in names}

    marks = ",".join("?" for _ in GENERIC_FERTILIZERS)
    cur.execute(f"DELETE FROM catalog_aliases WHERE catalog_id IN (SELECT id FROM catalog_entries WHERE division = ? AND by_product IN ({marks}))",
                (DIVISION, *GENERIC_FERTILIZERS))
    cur.execute(f"DELETE FROM catalog_entries WHERE division = ? AND by_product IN ({marks})", (DIVISION, *GENERIC_FERTILIZERS))
    cur.execute("SELECT coalesce(max(id), 0) FROM catalog_entries")
    next_id = cur.fetchone()[0] + 1
    for k, name in enumerate(names):
        cur.execute(
            "INSERT INTO catalog_entries (id, division, by_product, normalized_key) VALUES (?, ?, ?, ?)",
            (next_id + k, DIVISION, name, _key(name)),
        )
        # Exact match so a brand name never matches a different catalog entry.
        cur.execute("INSERT INTO catalog_aliases (catalog_id, raw_product, raw_by_product) VALUES (?, ?, ?)",
                    (next_id + k, DIVISION, name))

    cur.executemany(
        """INSERT INTO price_records (
               row_num, record_date, price_type, product, by_product,
               province, district, station, minimum, maximum, arrivals, arrivals_unit
           ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, ?)""",
        [
            (ROW_NUM_BASE + i, d, RATE_TYPE, DIVISION, name, PROVINCE, DISTRICT, STATION, lo, hi, unit)
            for i, (d, name, lo, hi, unit) in enumerate(rows)
        ],
    )

    with open(APP_BRANDS_JSON, "w", encoding="utf-8") as f:
        json.dump(
            {
                "_comment": "Generated by api/db/fertilizer.py from data/Fast uni.xlsx — do not edit by hand. "
                            "Fertilizer by-product -> company and brand (shown as 'Company' on the cards).",
                "byProduct": {n: info[n] for n in names},
            },
            f, ensure_ascii=False, indent=2,
        )
        f.write("\n")

    print(f"Fertilizer: {len(rows)} prices, {len(names)} brand products"
          + (f", skipped {skipped} incomplete rows" if skipped else "")
          + (f", left out {trimmed} outside {first}..{last}" if trimmed else ""))
    return len(rows)
